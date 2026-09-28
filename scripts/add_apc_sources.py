"""Populate the Sources tab of AP1000_Domestic_Content_Analysis.xlsx with one
hyperlinked row per source cited on the APCs tab, following the user's row-4 pattern
(Source | What It Provides | Source For)."""
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
import copy

FN = "AP1000_Domestic_Content_Analysis.xlsx"
try:
    f = open(FN, "r+b"); f.close()
except PermissionError:
    raise SystemExit("LOCKED - close the file in Excel first")

wb = load_workbook(FN)
ws = wb["Sources"]

b4, c4, d4 = ws["B4"], ws["C4"], ws["D4"]
link_font = copy.copy(b4.font) if b4.font else Font(color="0563C1", underline="single")
body_font = copy.copy(c4.font)
tag_font = copy.copy(d4.font)
wrap = Alignment(vertical="top", wrap_text=True)
thin = Border(bottom=Side(style="thin", color="D9D9D9"))

# fix truncated accession number + tighten description on the existing Cives row
b4.value = "NRC IR 99901419/2012-201 (ML13042A397)"
c4.value = ("Cives Steel Co., Southern Division, Thomasville, GA - fabrication of concrete "
            "embedments for Vogtle 3&4 and V.C. Summer 2&3 (inspection incl. Notice of Violation)")

NRC = "https://www.nrc.gov/docs/"
PSC = "https://psc.ga.gov/search/facts-document/?documentId="
rows = [
 ("IRS Notice 2023-38 (May 2023)", "https://www.irs.gov/pub/irs-drop/n-23-38.pdf",
  "Categorization framework and Table 2 precedents: rebar rows (Steel/Iron); offshore-wind monopile/transition piece = Manufactured Product (basis for CV/modules/piping positions); equipment = MP across all technologies",
  "APC #1, 6, 7, 8, 9-31 (framework)"),
 ("IRS Notice 2024-41 (May 2024)", "https://www.irs.gov/pub/irs-drop/n-24-41.pdf",
  "Hydropower Table 2 rows: 'embedded structure parts, foundation plates and anchors' and 'powerhouse structure' = Steel/Iron; GSU transformer = Manufactured Product",
  "APC #2, 3, 4, 22"),
 ("NRC IR 99901433/2017-201 (ML17355A050)", NRC + "ML1735/ML17355A050.pdf",
  "Newport News Industrial, Newport News, VA - shield building air inlet and tension ring panels for Vogtle 3&4", "APC #4"),
 ("GA PSC VCM 9/10 report (docket 29849, doc 152164)", PSC + "152164",
  "Newport News Industrial 'in full production of the shield building structural panels'; CA04 module fabricated on site", "APC #4, 7"),
 ("NRC IR 99901395/2010-201 (ML102870167)", NRC + "ML1028/ML102870167.pdf",
  "IHI Corporation, Yokohama, Japan - containment vessel fabrication for Vogtle 3&4 (plates/heads)", "APC #6"),
 ("GA PSC VCM 7 report (docket 29849, doc 143954)", PSC + "143954",
  "CB&I on-site assembly of containment vessels; Doosan RV/SG fabrication status; Curtiss-Wright RCP testing in PA; Tioga/IBF RCL piping", "APC #6, 8, 9, 10, 11, 17"),
 ("Nuclear Engineering Int'l (Aug 9, 2010)", "https://www.neimagazine.com/news/ihi-ships-heavy-components-for-vogtle-project/",
  "'IHI ships heavy components for Vogtle project': CV bottom heads produced at IHI works, Yokohama; CB&I awarded IHI the CV contract Jan 2009", "APC #6"),
 ("Southern Co./Georgia Power press release (May 9, 2014)", "https://www.prnewswire.com/news-releases/18-million-pound-containment-vessel-bottom-head-placed-at-vogtle-unit-4-258625111.html",
  "Unit 4 CV bottom head 'fabricated on site by CB&I'", "APC #6"),
 ("NRC IR 99901401/2012-201 (ML12279A119)", NRC + "ML1227/ML12279A119.pdf",
  "Shaw Modular Solutions, Lake Charles, LA - structural equipment modules for Vogtle and V.C. Summer", "APC #7"),
 ("GA PSC VCM 11 report (docket 29849, doc 154916)", PSC + "154916",
  "Unit 4 module re-sourcing: Oregon Iron Works, SMCI, IHI and Toshiba (CA01 fabricated in Japan); Mangiarotti pressurizer shipped July 2014", "APC #7, 12"),
 ("Toshiba, 5th MDEP Conference presentation (OECD-NEA, 2023)", "https://www.oecd-nea.org/mdep/events/conf-2023/presentations/S3/4.Toshiba.pdf",
  "Delivery record: Vogtle 3&4 STG & associated equipment (TC6F-52, 2013-14); CA01 sub-modules for Vogtle 4 (2016); turbine-island scope incl. deaerator and feedwater heaters", "APC #7, 23, 25"),
 ("NRC IR 99901428/2013-201 (ML13200A220)", NRC + "ML1320/ML13200A220.pdf",
  "IBF S.p.A., San Nicolo, Italy - reactor coolant loop seamless piping for the AP1000 fleet", "APC #8, 17"),
 ("NRC IR 99901432/2013-201 (ML13263A411)", NRC + "ML1326/ML13263A411.pdf",
  "CB&I Laurens, SC (formerly B.F. Shaw) - safety-related piping and piping modules", "APC #8"),
 ("NRC IR 99901373/2014-201 (ML14260A350)", NRC + "ML1426/ML14260A350.pdf",
  "Doosan, Changwon, South Korea - reactor vessels, closure heads, steam generators, RCP casings; Vogtle 3 RV nozzle weld corrective actions", "APC #9, 10, 11"),
 ("GA PSC VCM 4 report (docket 29849, doc 134428)", PSC + "134428",
  "Heavy forgings shipped from Japan Steel Works (Muroran, Japan) to Doosan (Changwon); Shaw Modular Solutions Lake Charles module facility", "APC #9, 7"),
 ("POWER Magazine (Nov 2016)", "https://www.powermag.com/reactor-vessel-placed-inside-vogtle-unit-3/",
  "Vogtle 3 reactor vessel 'fabricated by Doosan Heavy Industries in South Korea'", "APC #9"),
 ("World Nuclear News (Aug 17, 2017)", "https://world-nuclear-news.org/Articles/First-steam-generator-in-place-at-Vogtle",
  "All four Vogtle 3&4 steam generators 'fabricated in South Korea'", "APC #10"),
 ("NRC IR 99901383/2016-201 (ML16350A067)", NRC + "ML1635/ML16350A067.pdf",
  "Curtiss-Wright EMD, Cheswick, PA - design, manufacture, test and delivery of AP1000 reactor coolant pumps", "APC #11"),
 ("Georgia Power press release (Apr 18, 2016)", "https://www.prnewswire.com/news-releases/first-ap1000-reactor-coolant-pump-delivered-to-vogtle-expansion-300253080.html",
  "First RCP delivered to Vogtle from Curtiss-Wright, Cheswick, PA", "APC #11"),
 ("NRC IR 99901416/2012-201 (ML12320A661)", NRC + "ML1232/ML12320A661.pdf",
  "Mangiarotti, Monfalcone, Italy - accumulators, core makeup tanks, PRHR heat exchanger, pressurizer", "APC #12, 15"),
 ("GA PSC VCM 8 report (docket 29849, doc 146550)", PSC + "146550",
  "Mangiarotti completion and hydro-testing of Unit 3 accumulator tanks", "APC #15"),
 ("NRC IR 99901392/2014-201 (ML14328A138)", NRC + "ML1432/ML14328A138.pdf",
  "Westinghouse Newington Operations, Newington, NH - major supplier of CRDMs and reactor vessel internals for AP1000", "APC #13"),
 ("NRC IR 99901394/2015-201 (ML15132A142)", NRC + "ML1513/ML15132A142.pdf",
  "Premier Technology, Blackfoot, ID - integrated head package and RV internals lifting rig", "APC #14"),
 ("NRC IR 99900080/2012-201 (ML12158A154)", NRC + "ML1215/ML12158A154.pdf",
  "SPX Copes-Vulcan, PA - AP1000 squib valve design and manufacture (Vogtle / V.C. Summer)", "APC #16"),
 ("NRC IR 99901431/2014-201 (ML14073A652)", NRC + "ML1407/ML14073A652.pdf",
  "Pentair, Mansfield, MA - PV-62 pressurizer safety valves shipped for AP1000", "APC #16"),
 ("NRC IR 99901468/2016-201 (ML16160A176)", NRC + "ML1616/ML16160A176.pdf",
  "GE Oil & Gas (Consolidated), Pineville, LA - PV-65 main steam safety valves", "APC #16"),
 ("NRC IR 99901377/2012-201 (ML12306A385)", NRC + "ML1230/ML12306A385.pdf",
  "Enertech, Brea, CA - PXS nozzle check valve qualification", "APC #16"),
 ("Power Engineering (Jul 25, 2006)", "https://www.power-eng.com/nuclear/westinghouse-completes-purchase-of-par-nuclear/",
  "Westinghouse PaR Nuclear, Shoreview, MN - AP1000 fuel handling equipment and outage-critical cranes product line", "APC #18"),
 ("Holtec International news release (Feb 22, 2013)", "https://holtecinternational.com/2013/02/22/three-decades-of-relentless-drive-for-technology-improvement-cements-holtecs-role-as-the-worlds-preeminent-fuel-rack-supplier/",
  "High-density rack design/licensing for AP1000; rack module supply for Vogtle 3&4 and V.C. Summer 2&3", "APC #19"),
 ("NRC WEC PMS inspection (ML17123A085)", NRC + "ML1712/ML17123A085.pdf",
  "Westinghouse - design, implementation and testing of the PMS for Vogtle 3&4 (Warrendale/Cranberry, PA)", "APC #20"),
 ("NRC IR 99900265/2013-201 (ML13164A351)", NRC + "ML1316/ML13164A351.pdf",
  "General Atomics ESI, San Diego, CA - AP1000 radiation monitoring system", "APC #20"),
 ("NRC IR 99901435/2013-201 (ML14058A705)", NRC + "ML1405/ML14058A705.pdf",
  "EnerSys, Hays, KS - Class 1E batteries, qualifying for AP1000", "APC #21"),
 ("NRC RSCC inspection (ML14164A502)", NRC + "ML1416/ML14164A502.pdf",
  "RSCC Wire & Cable, East Granby, CT - Class 1E cable, AP1000 qualification in progress", "APC #21"),
 ("Toshiba press release (Dec 1, 2011)", "https://www.global.toshiba/ww/news/corporate/2011/12/pr0101.html",
  "Vogtle 3 condenser manufactured with BHI Co. Ltd at Sacheon, South Korea; Toshiba supplying Unit 4 condenser as well", "APC #24"),
 ("NRC IR 99901369/2013-201 (ML13119A154)", NRC + "ML1311/ML13119A154.pdf",
  "Flowserve Pump Division, Vernon, CA - safety-related pumps for AP1000", "APC #29"),
 ("Westinghouse press release (Oct 14, 2022)", "https://www.businesswire.com/news/home/20221014005260/en/Westinghouse-Congratulates-Partners-on-Fuel-Load-for-Vogtle-Unit-3",
  "Vogtle 3 fuel assemblies manufactured at the Columbia Fuel Fabrication Facility, SC", "APC #31"),
 ("No public source identified", "",
  "Vogtle vendors never publicly named: main step-up transformers, standby diesel generators, polar crane fabricator (PaR probable), pool/canal liner fabricator", "APC #5, 22, 26 (gap noted)"),
]

r = 5
for name, url, what, forwhat in rows:
    cb = ws.cell(row=r, column=2, value=name)
    if url:
        cb.hyperlink = url
        cb.font = link_font
    else:
        cb.font = body_font
    cc = ws.cell(row=r, column=3, value=what); cc.font = body_font
    cd = ws.cell(row=r, column=4, value=forwhat); cd.font = tag_font
    for c in (cb, cc, cd):
        c.alignment = wrap
        c.border = thin
    r += 1
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 90
ws.column_dimensions["D"].width = 26
wb.save(FN)
print(f"Sources tab: wrote {len(rows)} rows (rows 5-{r-1}); fixed B4 accession number")
